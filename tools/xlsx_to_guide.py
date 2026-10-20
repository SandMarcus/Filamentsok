"""Gjør masterdatabladet (Excel) om til data/guide.json, som guidesiden leser.

Bruk (krever openpyxl, bare på din egen maskin, ikke i GitHub Actions):
    pip install openpyxl
    python tools/xlsx_to_guide.py [data/guide/master.xlsx] [data/guide.json]

Endre arket, kjør kommandoen, og commit både arket og guide.json.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import guide_no as N  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SITE_KEY = {"PLA+ / Pro": "PLA+", "PA (Nylon)": "PA"}
SITE_KEYS = ["ABS", "ABS+", "APLA", "ASA", "ASA+", "BVOH", "CoPA", "CPE", "HIPS", "HT-PLA", "MABS", "PA6", "PA12", "PA612", "PA",
             "PAHT", "PBT", "PC", "PC-ABS", "PC-PBT", "PCL", "PCTG", "PE", "PEBA", "PEEK", "PEI", "PEKK", "PET", "PETG", "PETG+",
             "PHA", "PIPG", "PLA", "PLA+", "PMMA", "PP", "PPA", "PPS", "PPSU", "PSU", "PVA", "PVB", "PVDF", "SAN", "SBS", "SMP",
             "TPE", "TPR", "TPU"]


def lv(v, where):
    v = "" if v is None else str(v).strip()
    if v in N.LEVEL:
        return N.LEVEL[v]
    raise SystemExit(f"Ukjent verdi «{v}» i {where}. Legg den til i tools/guide_no.py (LEVEL).")


def mp(table, v, where):
    v = "" if v is None else str(v).strip()
    if v in table:
        return table[v]
    raise SystemExit(f"Ukjent verdi «{v}» i {where}. Legg den til i tools/guide_no.py.")


def temp(v, no=None, kind=""):
    """Temperaturer: norsk tekst hvis den finnes, ellers arkets verdi med °C hvis det er rene tall."""
    if no:
        return no
    v = "" if v is None else str(v).strip()
    if v == "Product-specific":
        return "Produktspesifikt"
    if v == "Manufacturer-specific":
        return "Produsentspesifikt"
    if re.fullmatch(r"[\d–\- ]+", v):
        return v.replace("-", "–") + " °C"
    return v


def sources(v):
    v = str(v or "")
    for a, b in [("Manufacturer TDS", "produsentens datablad"), ("manufacturer TDS", "produsentens datablad"),
                 ("Product TDS required", "produktets datablad kreves"), ("Manufacturer-specific", "produsentspesifikt"),
                 ("guidance", "veiledning"), ("guide", "guide")]:
        v = v.replace(a, b)
    return v[0].upper() + v[1:] if v else v


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/guide/master.xlsx"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data/guide.json"
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["Master Datasheet"]
    hdr = [c.value for c in ws[1]]
    refs = {}
    wr = wb["Manufacturer vs General Data"]
    for r in wr.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        item = {"s": r[2], "noz": r[3], "bed": r[4], "ch": r[5], "hdt": r[6], "tg": r[7], "sh": r[8], "ts": r[9], "fm": r[10],
                "dry": r[11], "n": N.REF_NOTE.get(str(r[12]))}
        if item["n"] is None:
            raise SystemExit(f"Mangler oversettelse av referansenotat: {r[12]!r} (tools/guide_no.py, REF_NOTE)")
        for k in N.REF_CLASS_KEYS[r[0]]:
            refs.setdefault(k, []).append(item)
    mats = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        d = dict(zip(hdr, row))
        name = d["Material"]
        if name not in N.NO:
            raise SystemExit(f"Mangler norsk tekst for {name} i tools/guide_no.py (NO).")
        t = N.NO[name]
        key = SITE_KEY.get(name, name)
        seen.add(key)
        w = f"{name}"
        L = lambda col: lv(d[col], f"{w} / {col}")
        mats.append({
            "id": re.sub(r"[^a-z0-9]+", "-", name.lower().replace("+", "plus")).strip("-"),
            "abbr": name,
            "key": key,
            "tech": N.NAME.get(d["Chemical / Full Name"], d["Chemical / Full Name"]),
            "fam": mp(N.FAMILY, d["Family"], f"{w} / Family"),
            "gen": {
                "diff": L("Print Difficulty"), "heat": L("Heat Resistance"), "moist": L("Moisture Sensitivity"),
                "warp": L("Warping Risk"), "tough": L("Impact/Toughness"), "flex": L("Flexibility"),
                "chem": L("Chemical Resistance"), "uv": L("UV Resistance"), "fat": L("Fatigue Resistance"),
                "encl": L("Enclosure"), "dry": L("Drybox"), "hard": L("Hardened Nozzle"),
                "uses": t["uses"], "avoid": t["avoid"],
            },
            "spec": {
                "noz": temp(d["Nozzle Temp °C"], t.get("noz")), "bed": temp(d["Bed Temp °C"], t.get("bed")),
                "cool": L("Cooling"), "hw": mp(N.HARDWARE, d["Typical Hardware"], f"{w} / Typical Hardware"),
                "drying": t["dry"], "str": t["str"], "stiff": L("Stiffness"), "impact": L("Impact Resistance"),
                "layer": L("Layer Adhesion"), "dim": L("Dimensional Stability"), "abr": L("Abrasion Resistance"),
                "elast": L("Elasticity"), "water": L("Water Resistance"), "uvw": L("UV / Weather"),
                "heatr": L("Heat Resistance Rating"), "chemr": L("Chemical Resistance Rating"),
                "fatr": L("Fatigue Resistance Rating"), "flexr": L("Flexibility Rating"), "sol": L("Soluble"),
                "food": mp(N.FOOD, d["Food Contact (FDA / EU 10/2011)"], f"{w} / Food"),
                "med": mp(N.MED, d["Medical / Biocompatibility (ISO 10993 / USP Class VI)"], f"{w} / Medical"),
                "flame": mp(N.FLAME, d["Flame Retardancy (UL94 Rating)"], f"{w} / UL94"),
                "notes": t["notes"], "src": sources(d["Source Family"]),
                "conf": mp(N.CONF, d["Data Confidence"], f"{w} / Confidence"),
            },
            "refs": refs.get(name, []),
        })
    missing = [k for k in SITE_KEYS if k not in seen]
    src_rows = [{"n": r[0], "u": r[1], "use": r[2]} for r in wb["Sources"].iter_rows(min_row=2, values_only=True) if r[0]]
    body = {"v": 1, "materials": mats, "sources": src_rows, "missing": missing}
    out.write_text(json.dumps(body, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(mats)} materialer skrevet til {out}")
    if missing:
        print("Finnes på nettsiden, men ikke i arket:", ", ".join(missing))


if __name__ == "__main__":
    main()
